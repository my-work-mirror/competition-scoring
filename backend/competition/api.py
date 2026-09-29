"""大赛评分接口：评委用 API Key 登录打分，管理员维护维度一与现场开关。"""

from __future__ import annotations

import logging
from io import BytesIO
from pathlib import Path
from typing import Callable

from flask import Blueprint, jsonify, request, send_file

from ..services import key_auth
from . import scoring
from .roster import (
    CompetitionConfig,
    CompetitionConfigError,
    Person,
    load_competition_config,
)
from .store import CompetitionStore

logger = logging.getLogger(__name__)

APP_SLUG = "competition-scoring"
API_PREFIX = f"/api/ai-apps/{APP_SLUG}"
ROLE_ADMIN = "admin"
ROLE_JUDGE = "judge"


def create_competition_blueprint(
    *,
    config_file: Path,
    db_file: Path,
    resolve_identity: Callable[[str], dict | None],
    token_secret: str,
) -> Blueprint:
    blueprint = Blueprint("competition_scoring", __name__)
    caches: dict[str, object] = {}
    # 启动即校验，配置写错时在注册阶段就暴露，而不是等到现场打分。
    load_competition_config(config_file)

    def config() -> CompetitionConfig:
        """按文件 mtime 热加载，现场改名册或队伍无需重启服务。"""
        cached = caches.get("config")
        stamp = config_file.stat().st_mtime if config_file.exists() else 0
        if cached is None or caches.get("config_stamp") != stamp:
            cached = load_competition_config(config_file)
            caches["config"] = cached
            caches["config_stamp"] = stamp
        return cached  # type: ignore[return-value]

    def store() -> CompetitionStore:
        if "store" not in caches:
            caches["store"] = CompetitionStore(db_file)
        return caches["store"]  # type: ignore[return-value]

    # ---------- 认证 ----------

    def bearer_token() -> str:
        header = request.headers.get("Authorization", "")
        if header.startswith("Bearer "):
            return header[len("Bearer ") :].strip()
        payload = request.get_json(silent=True) or {}
        return str(payload.get("token", "")).strip()

    def authenticate(required_role: str | None = None):
        """校验 token 并返回 (身份, 错误响应)。"""
        info = key_auth.verify_token(bearer_token(), token_secret)
        if info is None:
            return None, (jsonify({"error": "登录状态已过期，请重新用 API Key 登录"}), 401)
        role = info.get("scope", "")
        if role not in {ROLE_ADMIN, ROLE_JUDGE}:
            return None, (jsonify({"error": "凭证与评分系统不匹配"}), 403)
        employee_id = info.get("employee_id", "")
        resolved = config().resolve_person(employee_id=employee_id)
        if resolved is None:
            return None, (jsonify({"error": "该账号已不在评委名册中"}), 403)
        person, current_role = resolved
        if required_role == ROLE_ADMIN and current_role != ROLE_ADMIN:
            return None, (jsonify({"error": "需要管理员权限"}), 403)
        return {"person": person, "role": current_role}, None

    @blueprint.post(f"{API_PREFIX}/auth")
    def competition_auth():
        payload = request.get_json(silent=True) or {}
        api_key = str(payload.get("api_key", "")).strip()
        if not api_key:
            return jsonify({"error": "请输入 API Key"}), 400

        try:
            settings_config = config()
        except CompetitionConfigError as exc:
            logger.exception("competition config invalid")
            return jsonify({"error": str(exc)}), 503

        key_hash = key_auth.hash_api_key(api_key)
        employee_id = ""
        email = ""
        try:
            record = resolve_identity(key_hash)
        except RuntimeError as exc:
            logger.warning("failed to resolve competition key: %s", exc)
            return jsonify({"error": "身份服务暂时不可用，请稍后重试"}), 503
        if record is not None:
            employee_id = record.get("employee_id", "")
            email = record.get("email", "") or ""

        resolved = settings_config.resolve_person(
            employee_id=employee_id, email=email, key_hash=key_hash
        )
        if resolved is None:
            return jsonify(
                {
                    "code": "not_a_judge",
                    "error": "该 API Key 不在评委名册中，请联系 AI应用部",
                }
            ), 403
        person, role = resolved
        if not person.employee_id:
            return jsonify({"error": "评委名册缺少 employee_id，无法建立稳定身份"}), 503

        token, expires_at = key_auth.issue_token(
            key_hash,
            token_secret,
            scope=role,
            employee_id=person.employee_id,
        )
        return jsonify(
            {
                "token": token,
                "expires_at": expires_at,
                "role": role,
                "judge": {"employee_id": person.employee_id, "name": person.name},
                **state_payload(person, role),
            }
        )

    # ---------- 公共视图 ----------

    def team_payload(team) -> dict:
        return {
            "id": team.id,
            "display_order": team.display_order,
            "name": team.name,
            "contest_label": team.contest_label,
            "department_ids": list(team.department_ids),
        }

    def dimension_two_summary(competition_id: str) -> dict[str, dict]:
        current = config()
        active_judge_ids = {judge.employee_id for judge in current.judges}
        grouped: dict[str, list[dict]] = {team.id: [] for team in current.teams}
        for row in store().all_judge_scores(competition_id):
            if row["team_id"] in grouped and row["judge_employee_id"] in active_judge_ids:
                grouped[row["team_id"]].append(row)
        summary: dict[str, dict] = {}
        for team_id, rows in grouped.items():
            result = scoring.trimmed_average(
                [row["total"] for row in rows],
                trim_high=current.trim_high,
                trim_low=current.trim_low,
            )
            summary[team_id] = {
                "total": result.total,
                "judge_count": result.judge_count,
                "counted_count": result.counted_count,
                "dropped_high": result.dropped_high,
                "dropped_low": result.dropped_low,
                "trimmed": result.trimmed,
            }
        return summary

    def leaderboard(competition_id: str) -> list[dict]:
        current = config()
        dimension_one = store().dimension_one_rows(competition_id)
        dimension_two = dimension_two_summary(competition_id)
        rows = []
        for team in current.teams:
            first = dimension_one.get(team.id) or {}
            second = dimension_two.get(team.id) or {}
            rows.append(
                {
                    "team_id": team.id,
                    "name": team.name,
                    "contest_label": team.contest_label,
                    "display_order": team.display_order,
                    "dimension_one": scoring.round_half_up(float(first.get("total", 0) or 0)),
                    "dimension_two": float(second.get("total", 0) or 0),
                    "judge_count": second.get("judge_count", 0),
                    "counted_count": second.get("counted_count", 0),
                    "trimmed": second.get("trimmed", False),
                }
            )
        return scoring.rank_teams(rows, list(current.awards))

    def state_payload(person: Person, role: str) -> dict:
        current = config()
        settings = store().settings(current.id)
        dimension_one = store().dimension_one_rows(current.id)
        my_scores = store().judge_scores_for_judge(current.id, person.employee_id)
        published = settings.get("results_published") == "1"
        payload = {
            "competition": {
                "id": current.id,
                "name": current.name,
                "organizer": current.organizer,
                "dimension_one_label": current.dimension_one_label,
                "dimension_two_label": current.dimension_two_label,
                "dimension_one_max": current.dimension_one_max,
                "dimension_two_max": current.dimension_two_max,
                "score_step": current.score_step,
                "trim_high": current.trim_high,
                "trim_low": current.trim_low,
            },
            "items": [
                {"key": item.key, "name": item.name, "max": item.max, "criteria": item.criteria}
                for item in current.items
            ],
            "awards": list(current.awards),
            "settings": {
                "scoring_open": settings.get("scoring_open") == "1",
                "results_published": published,
            },
            "teams": [
                {
                    **team_payload(team),
                    "dimension_one": scoring.round_half_up(
                        float((dimension_one.get(team.id) or {}).get("total", 0) or 0)
                    ),
                    "dimension_one_detail": {
                        "training_count": (dimension_one.get(team.id) or {}).get(
                            "training_count", 0
                        ),
                        "training_score": (dimension_one.get(team.id) or {}).get(
                            "training_score", 0
                        ),
                        "audience_total": (dimension_one.get(team.id) or {}).get(
                            "audience_total", 0
                        ),
                        "signin_score": (dimension_one.get(team.id) or {}).get("signin_score", 0),
                        "source": (dimension_one.get(team.id) or {}).get("source", ""),
                        "note": (dimension_one.get(team.id) or {}).get("note", ""),
                    },
                    "my_score": my_scores.get(team.id),
                }
                for team in current.teams
            ],
            "judges": [
                {"employee_id": judge.employee_id, "name": judge.name} for judge in current.judges
            ],
        }
        if role == ROLE_ADMIN or published:
            payload["leaderboard"] = leaderboard(current.id)
        return payload

    @blueprint.get(f"{API_PREFIX}/state")
    def competition_state():
        identity, error = authenticate()
        if error is not None:
            return error
        return jsonify(
            {
                "role": identity["role"],
                "judge": {
                    "employee_id": identity["person"].employee_id,
                    "name": identity["person"].name,
                },
                **state_payload(identity["person"], identity["role"]),
            }
        )

    # ---------- 评委打分 ----------

    @blueprint.post(f"{API_PREFIX}/scores")
    def submit_score():
        identity, error = authenticate()
        if error is not None:
            return error
        current = config()
        settings = store().settings(current.id)
        if settings.get("scoring_open") != "1" and identity["role"] != ROLE_ADMIN:
            return jsonify({"error": "打分已截止，如需修改请联系 AI应用部"}), 409

        payload = request.get_json(silent=True) or {}
        team_id = str(payload.get("team_id", "")).strip()
        team = current.team(team_id)
        if team is None:
            return jsonify({"error": "参赛部门不存在"}), 404

        raw_items = payload.get("items")
        if not isinstance(raw_items, dict):
            return jsonify({"error": "items 必须是评分项对象"}), 400

        items: dict[str, float] = {}
        for item in current.items:
            if item.key not in raw_items:
                return jsonify({"error": f"缺少评分项：{item.name}"}), 400
            try:
                value = float(raw_items[item.key])
            except (TypeError, ValueError):
                return jsonify({"error": f"{item.name} 的分值不是数字"}), 400
            if value < 0 or value > item.max:
                return jsonify({"error": f"{item.name} 应在 0 - {item.max:g} 分之间"}), 400
            step = current.score_step
            if step > 0 and abs(round(value / step) - value / step) > 1e-9:
                return jsonify({"error": f"{item.name} 需为 {step:g} 的整数倍"}), 400
            items[item.key] = value

        total = scoring.round_half_up(sum(items.values()))
        comment = str(payload.get("comment", "")).strip()[:2000]
        row = store().put_judge_score(
            current.id,
            team_id,
            identity["person"].employee_id,
            identity["person"].name,
            items,
            total,
            comment,
        )
        return jsonify({"score": row})

    @blueprint.get(f"{API_PREFIX}/my-scores")
    def my_scores():
        identity, error = authenticate()
        if error is not None:
            return error
        current = config()
        return jsonify(
            {"scores": store().judge_scores_for_judge(current.id, identity["person"].employee_id)}
        )

    @blueprint.get(f"{API_PREFIX}/results")
    def results():
        identity, error = authenticate()
        if error is not None:
            return error
        current = config()
        settings = store().settings(current.id)
        if identity["role"] != ROLE_ADMIN and settings.get("results_published") != "1":
            return jsonify({"error": "结果尚未公布"}), 403
        return jsonify({"leaderboard": leaderboard(current.id)})

    # ---------- 管理员 ----------

    @blueprint.get(f"{API_PREFIX}/admin/progress")
    def admin_progress():
        identity, error = authenticate(ROLE_ADMIN)
        if error is not None:
            return error
        current = config()
        rows = store().all_judge_scores(current.id)
        by_team: dict[str, dict[str, dict]] = {team.id: {} for team in current.teams}
        for row in rows:
            if row["team_id"] in by_team:
                by_team[row["team_id"]][row["judge_employee_id"]] = row
        matrix = []
        for team in current.teams:
            submitted = by_team.get(team.id, {})
            matrix.append(
                {
                    "team_id": team.id,
                    "name": team.name,
                    "contest_label": team.contest_label,
                    "submitted_count": len(submitted),
                    "judges": [
                        {
                            "employee_id": judge.employee_id,
                            "name": judge.name,
                            "submitted": judge.employee_id in submitted,
                            "total": submitted.get(judge.employee_id, {}).get("total"),
                            "items": submitted.get(judge.employee_id, {}).get("items"),
                            "comment": submitted.get(judge.employee_id, {}).get("comment", ""),
                            "updated_at": submitted.get(judge.employee_id, {}).get("updated_at"),
                        }
                        for judge in current.judges
                    ],
                }
            )
        return jsonify({"matrix": matrix, "judge_total": len(current.judges)})

    @blueprint.post(f"{API_PREFIX}/admin/settings")
    def admin_settings():
        identity, error = authenticate(ROLE_ADMIN)
        if error is not None:
            return error
        current = config()
        payload = request.get_json(silent=True) or {}
        allowed = {"scoring_open", "results_published"}
        updates = {key: payload[key] for key in allowed if key in payload}
        if not updates:
            return jsonify({"error": "没有可更新的开关"}), 400
        settings = store().settings(current.id)
        for key, value in updates.items():
            settings = store().set_setting(
                current.id,
                key,
                "1" if bool(value) else "0",
                actor_employee_id=identity["person"].employee_id,
                actor_name=identity["person"].name,
            )
        return jsonify(
            {
                "settings": {
                    "scoring_open": settings.get("scoring_open") == "1",
                    "results_published": settings.get("results_published") == "1",
                }
            }
        )

    def apply_dimension_one(entry: dict, actor: Person) -> dict:
        """把一条维度一录入落库。总分优先用显式 total，否则按规则推导。"""
        current = config()
        team_id = str(entry.get("team_id", "")).strip()
        team = current.team(team_id)
        if team is None:
            raise ValueError(f"参赛部门不存在：{team_id}")
        training_count = int(entry.get("training_count", 0) or 0)
        audience_total = int(entry.get("audience_total", 0) or 0)
        if training_count < 0 or audience_total < 0:
            raise ValueError("培训次数与覆盖人次不能为负数")
        derived = scoring.dimension_one(
            training_count,
            audience_total,
            training_max=current.training_count_max,
            signin_max=current.signin_max,
            full_headcount=current.signin_full_headcount,
        )
        payload = {
            "training_count": derived.training_count,
            "training_score": derived.training_score,
            "audience_total": derived.audience_total,
            "signin_score": derived.signin_score,
            "total": derived.total,
            "source": str(entry.get("source", "manual")),
            "note": str(entry.get("note", "")).strip()[:500],
        }
        if entry.get("total") not in (None, ""):
            override = float(entry["total"])
            if override < 0 or override > current.dimension_one_max:
                raise ValueError(f"维度一总分应在 0 - {current.dimension_one_max:g} 之间")
            payload["total"] = scoring.round_half_up(override)
            payload["source"] = "override"
        return store().put_dimension_one(
            current.id,
            team_id,
            payload,
            actor_employee_id=actor.employee_id,
            actor_name=actor.name,
        )

    @blueprint.post(f"{API_PREFIX}/admin/dimension-one")
    def admin_dimension_one():
        identity, error = authenticate(ROLE_ADMIN)
        if error is not None:
            return error
        payload = request.get_json(silent=True) or {}
        entries = payload.get("entries")
        if entries is None:
            entries = [payload]
        if not isinstance(entries, list) or not entries:
            return jsonify({"error": "entries 必须是非空数组"}), 400
        saved = []
        try:
            for entry in entries:
                saved.append(apply_dimension_one(entry, identity["person"]))
        except (ValueError, TypeError) as exc:
            return jsonify({"error": str(exc)}), 400
        return jsonify({"rows": saved})

    @blueprint.post(f"{API_PREFIX}/admin/dimension-one/import")
    def admin_dimension_one_import():
        identity, error = authenticate(ROLE_ADMIN)
        if error is not None:
            return error
        upload = request.files.get("file")
        if upload is None or not upload.filename:
            return jsonify({"error": "请选择要导入的 Excel 文件"}), 400
        try:
            parsed = parse_training_workbook(upload.stream.read(), config())
        except ValueError as exc:
            return jsonify({"error": str(exc)}), 400

        saved = []
        for entry in parsed["matched"]:
            saved.append(apply_dimension_one({**entry, "source": "excel"}, identity["person"]))
        return jsonify({"rows": saved, "unmatched": parsed["unmatched"], "sheet": parsed["sheet"]})

    @blueprint.post(f"{API_PREFIX}/admin/reset-scores")
    def admin_reset_scores():
        identity, error = authenticate(ROLE_ADMIN)
        if error is not None:
            return error
        current = config()
        payload = request.get_json(silent=True) or {}
        if payload.get("confirm") is not True:
            return jsonify({"error": "需要显式确认才能清空打分"}), 400
        settings = store().settings(current.id)
        if settings.get("results_published") == "1":
            return jsonify(
                {"error": "结果已公布，请先取消公布再清空，避免公示成绩被静默改写"}
            ), 409
        cleared = store().clear_judge_scores(
            current.id,
            actor_employee_id=identity["person"].employee_id,
            actor_name=identity["person"].name,
        )
        return jsonify({"cleared": cleared})

    @blueprint.get(f"{API_PREFIX}/admin/audit")
    def admin_audit():
        identity, error = authenticate(ROLE_ADMIN)
        if error is not None:
            return error
        limit = request.args.get("limit", "200")
        try:
            parsed_limit = int(limit)
        except ValueError:
            parsed_limit = 200
        return jsonify({"entries": store().audit_trail(config().id, parsed_limit)})

    @blueprint.get(f"{API_PREFIX}/admin/export")
    def admin_export():
        identity, error = authenticate(ROLE_ADMIN)
        if error is not None:
            return error
        current = config()
        content = render_results_workbook(
            current,
            leaderboard(current.id),
            store().all_judge_scores(current.id),
            store().dimension_one_rows(current.id),
        )
        return send_file(
            BytesIO(content),
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=f"{current.id}-results.xlsx",
        )

    return blueprint


def parse_training_workbook(content: bytes, current: CompetitionConfig) -> dict:
    """从培训统计工作簿中读出各部门培训次数与累计签到人次。"""
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:  # openpyxl 对坏文件抛出多种异常
        raise ValueError(f"无法读取 Excel 文件：{exc}") from exc

    sheet = None
    for candidate in workbook.worksheets:
        headers = [str(cell or "") for cell in next(candidate.iter_rows(values_only=True), ())]
        joined = "".join(headers)
        if "培训数量" in joined or "培训次数" in joined:
            sheet = candidate
            break
    if sheet is None:
        raise ValueError("未找到含“培训数量/培训次数”表头的工作表")

    header_row = [str(cell or "").strip() for cell in next(sheet.iter_rows(values_only=True), ())]

    def column_of(*keywords: str) -> int | None:
        for index, header in enumerate(header_row):
            if any(keyword in header for keyword in keywords):
                return index
        return None

    unit_column = column_of("单位", "部门") or 0
    count_column = column_of("培训数量", "培训次数")
    audience_column = column_of("累计听众", "听众人数", "覆盖人数", "人次")
    if count_column is None or audience_column is None:
        raise ValueError("表头缺少培训数量或累计听众人数列")

    matched: list[dict] = []
    unmatched: list[str] = []
    seen: set[str] = set()
    for row in sheet.iter_rows(min_row=2, values_only=True):
        label = str(row[unit_column] or "").strip() if unit_column < len(row) else ""
        if not label:
            continue
        team = current.team_by_label(label)
        if team is None:
            unmatched.append(label)
            continue
        if team.id in seen:
            continue
        seen.add(team.id)

        def cell(index: int) -> int:
            if index >= len(row) or row[index] in (None, ""):
                return 0
            try:
                return int(float(row[index]))
            except (TypeError, ValueError):
                return 0

        matched.append(
            {
                "team_id": team.id,
                "training_count": cell(count_column),
                "audience_total": cell(audience_column),
                "note": f"导入自 {sheet.title}（原始名称 {label}）",
            }
        )
    if not matched:
        raise ValueError("没有任何一行能匹配到参赛部门，请检查部门名称")
    return {"matched": matched, "unmatched": unmatched, "sheet": sheet.title}


def render_results_workbook(
    current: CompetitionConfig,
    ranked: list[dict],
    judge_scores: list[dict],
    dimension_one: dict[str, dict],
) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()

    overview = workbook.active
    overview.title = "总排名"
    overview.append(
        [
            "名次",
            "参赛部门",
            "通知简称",
            current.dimension_one_label,
            current.dimension_two_label,
            "最终得分",
            "奖项",
            "有效评委数",
            "计入评委数",
        ]
    )
    for row in ranked:
        overview.append(
            [
                row["rank"],
                row["name"],
                row["contest_label"],
                row["dimension_one"],
                row["dimension_two"],
                row["total"],
                row["award"],
                row["judge_count"],
                row["counted_count"],
            ]
        )

    detail = workbook.create_sheet("评委明细")
    detail.append(
        ["参赛部门", "评委", *[item.name for item in current.items], "小计", "评语", "修订次数"]
    )
    team_names = {team.id: team.name for team in current.teams}
    for row in judge_scores:
        items = row.get("items") or {}
        detail.append(
            [
                team_names.get(row["team_id"], row["team_id"]),
                row.get("judge_name", ""),
                *[items.get(item.key, "") for item in current.items],
                row.get("total", 0),
                row.get("comment", ""),
                row.get("revision", 1),
            ]
        )

    training = workbook.create_sheet(current.dimension_one_label)
    training.append(["参赛部门", "培训次数", "培训分数", "累计听众人次", "签到分数", "小计", "来源", "备注"])
    for team in current.teams:
        row = dimension_one.get(team.id) or {}
        training.append(
            [
                team.name,
                row.get("training_count", 0),
                row.get("training_score", 0),
                row.get("audience_total", 0),
                row.get("signin_score", 0),
                row.get("total", 0),
                row.get("source", ""),
                row.get("note", ""),
            ]
        )

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
